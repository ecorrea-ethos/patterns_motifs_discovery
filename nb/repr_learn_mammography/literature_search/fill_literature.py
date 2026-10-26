"""Fill Triage suggestion (M), Supported literature (O) and PMIDs (P) of outputs/final_table_mammography_k8_v2.xlsx and write the audit files.

usage (conda env ethos_llm, cwd = literature_search/):
    python fill_literature.py --check              # test the decisions against the abstracts, write nothing
    python fill_literature.py [--out PATH.xlsx]    # write the workbook (default: the table itself, in place)

The table lists only the NMF bundles c0-c7 (24 rows, Excel rows 2-25). PAT in pubmed.py holds the steps of each row; if cell 48 of the
notebook is re-run and the patterns change, update PAT, re-run round1.py / round2.py, fetch_abstracts.py, then KEEP below.
Inputs: round1.log / round2.log (what PubMed returned in this run), abstracts.json (efetch of every kept paper).
A row keeps a paper only if its title or abstract names at least min(2, number of distinct steps) of the row's steps (steps.py) and the
title fits the bundle's setting. PubMed cannot test order, so no row is rated Yes. A row with no kept paper is None / 'none found'.
Writes next to the workbook: *_literature_evidence.csv (per row and paper: steps found / not found), *_literature_search_log.csv.
"""
import argparse, csv, json, os, re
import openpyxl
from pubmed import PAT
from steps import STEP

HERE = os.path.dirname(os.path.abspath(__file__))
TABLE = os.path.join(HERE, '..', 'outputs', 'final_table_mammography_k8_v2.xlsx')
ap = argparse.ArgumentParser(); ap.add_argument('--out', default=TABLE); ap.add_argument('--check', action='store_true')
args = ap.parse_args()

# bundle (first row of its 3-row block) -> triage (draft, from the bundle's entities and name, not from the literature)
TRIAGE = {2: 'Routine', 5: 'Unclear', 8: 'Routine', 11: 'Routine', 14: 'Routine', 17: 'Routine', 20: 'Manual RUW', 23: 'Unclear'}
# table row -> PMIDs kept (all Partly: no paper tests the order of the steps). Rows missing here have no kept paper -> None.
KEEP = {2: ['23317023', '12133160', '9179135', '28671391'], 3: ['23317023', '12133160', '9179135', '28671391'],
        4: ['23317023', '12133160', '9179135', '33890501'],
        5: ['32352727'], 6: ['32352727'], 7: ['32352727'],
        11: ['23317023', '33711775', '21225474', '31331610'], 12: ['23317023', '33711775', '21225474', '31331610'],
        13: ['15963877', '21225474', '31331610'],
        14: ['35440339', '35743024', '39107936'], 15: ['34288808', '35743024', '39107936'], 16: ['35801218', '35743024', '39107936'],
        17: ['29373057', '34152897', '40325660'], 18: ['29373057', '28053200', '31728373', '40887656'],
        19: ['29373057', '34152897', '22020596'],
        20: ['17163275', '15190136', '12859302'], 21: ['17163275', '15190136', '12859302'], 22: ['15190136', '12859302', '17163275'],
        23: ['26700342', '26658934', '41105409', '37639247'], 24: ['26700342', '26658934', '41105409', '37639247'],
        25: ['37639247', '41105409', '26658934', '20800770']}
abstracts = json.load(open(os.path.join(HERE, 'abstracts.json')))
returned = open(os.path.join(HERE, 'round1.log')).read() + open(os.path.join(HERE, 'round2.log')).read()

evidence, problems = [], []
for row, pmids in KEEP.items():
    steps = list(dict.fromkeys(PAT[row])); need = min(2, len(steps))
    assert len(set(pmids)) == len(pmids)
    for pm in pmids:
        if f'{pm} ' not in returned:
            problems.append(f'row {row}: PMID {pm} was not returned by any search of this run')
        d = abstracts.get(pm)
        if not d:
            problems.append(f'row {row}: PMID {pm} has no fetched record'); continue
        text = d['title'] + ' ' + d['abstract']
        found = [t for t in steps if re.search(STEP[t], text, re.I)]
        basis = ('abstract' if d['abstract'] else 'title only') if len(found) >= need else 'FAILS'
        if basis == 'FAILS':
            problems.append(f'row {row}: PMID {pm} names only {found} of {steps}')
        evidence.append([row, pm, d['year'], d['journal'], d['title'], ' '.join(steps), ' '.join(found), ' '.join(t for t in steps if t not in found), basis])
if problems:
    raise SystemExit('\n'.join(problems))
n_none = [r for r in range(2, 26) if r not in KEEP]
print(f'{len(KEEP)} rows Partly, {len(n_none)} rows None {n_none}; {len({p for v in KEEP.values() for p in v})} distinct PMIDs, all returned in this run and all pass the step check')
if args.check:
    raise SystemExit(0)

wb = openpyxl.load_workbook(TABLE)
ws = wb.active
hdr = {c.value: c.column_letter for c in ws[1]}
colM, colO, colP = hdr['Triage suggestion (Manual RUW / Routine / Unclear)'], hdr['Supported literature (Yes / Partly / None)'], hdr['PMIDs']
assert (colM, colO, colP) == ('M', 'O', 'P')
assert all(ws[f'{c}{r}'].value is None for c in 'MOP' for r in range(2, 26)), 'the three columns are not empty: refusing to overwrite'
assert str(ws['A2'].value).startswith('c0 ') and str(ws['A23'].value).startswith('c7 '), 'the table must list only the bundles c0-c7 (re-run cell 48 of the v2 notebook)'
for r in range(2, 26):
    ws[f'O{r}'] = 'Partly' if r in KEEP else 'None'
    ws[f'P{r}'] = '; '.join(KEEP[r]) if r in KEEP else 'none found'
for r, t in TRIAGE.items():
    assert any(str(m) == f'M{r}:M{r + 2}' for m in ws.merged_cells.ranges), r
    ws[f'M{r}'] = t                                   # top-left cell of the merged M{r}:M{r+2}
note = next(r for r in range(26, ws.max_row + 1) if str(ws[f'A{r}'].value).startswith('Triage suggestion, Supported literature'))
IN_SITU = next((str(ws[f'A{r}'].value) for r in range(note, ws.max_row + 1) if str(ws[f'A{r}'].value).startswith('In-situ note')), '')
ws[f'A{note}'] = ("Triage suggestion (one per bundle, draft for underwriter review, from the bundle's entities and name, not from the literature): "
                  "Manual RUW = surgical / peri-operative breast cancer treatment shown (c6); Routine = care that screening, "
                  "check-ups or unrelated care explains; Unclear = enriched for cancer (OR 1.7 to 1.8) but the pattern is not specific to breast cancer (c1, c7). "
                  "Supported literature and PMIDs (filled 2026-10-19): PubMed search of at least 3 of the bundle's top-5 entities ANDed (unordered) or the steps of the typical order of care ANDed, "
                  "then the bundle's clinical setting (screening, staging labs, cardio-oncology, breast surgery anaesthesia, rehabilitation); a paper is kept only if its title fits the setting and its title or abstract names at least 2 of the row's steps "
                  "(1 when the row has a single distinct step; the 'chemistry panel' step also counts liver function tests). PubMed cannot test order, so no row is Yes; Partly = the steps are addressed in a more general or different setting; "
                  "None = no paper named the steps ('none found'); every PMID was returned by a search of that run and checked in PubMed. A clinician should confirm each Partly. "
                  "Evidence per row: final_table_mammography_k8_v2_literature_evidence.csv.")
wb.save(args.out)
stem = args.out.rsplit('.', 1)[0]
with open(stem + '_literature_evidence.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['excel_row', 'PMID', 'year', 'journal', 'title (as returned by PubMed)', 'row steps', 'steps found in title/abstract', 'steps not found', 'basis'])
    w.writerows(evidence)
with open(stem + '_literature_search_log.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['round', 'kind', 'key', 'query', 'pubmed_hits'])
    for x in json.load(open(os.path.join(HERE, 'lit_results.json'))): w.writerow([1, x['kind'], x['key'] + ' ' + x['concepts'], x['query'], x['count']])
    for x in json.load(open(os.path.join(HERE, 'lit2_results.json'))): w.writerow([2, 'setting', x['label'], x['query'], x['count']])
print('wrote', args.out)
