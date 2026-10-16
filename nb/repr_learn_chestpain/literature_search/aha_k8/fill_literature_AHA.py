"""Fill Triage suggestion (M), Supported literature (O) and PMIDs (P) of final_table_chestpain_k8_AHA.xlsx and write the audit files.

usage (conda env ethos_llm, cwd = literature_search/aha_k8/):
    python fill_literature_AHA.py --check                 # only test the decisions against the abstracts, write nothing
    python fill_literature_AHA.py --out PATH.xlsx         # write the workbook to PATH (default: the table itself, in place)

Inputs: round1_AHA.log / round2_AHA.log (what PubMed returned in this run), abstracts_AHA.json (efetch of every kept paper).
Every PMID below was returned by a search of this run and read (title + abstract). A row keeps a paper only if its title or abstract
names at least 2 of the row's steps, or it is a guideline / consensus statement on chest pain evaluation kept on its title (BASIS_TITLE).
PubMed cannot test order, so no row is rated Yes.
Writes next to the workbook: *_literature_evidence.csv (per row and paper: steps found / not found), *_literature_search_log.csv.
"""
import argparse, csv, json, os, re
import openpyxl
from pubmed_aha import PAT

HERE = os.path.dirname(os.path.abspath(__file__))
TABLE = os.path.join(HERE, '..', '..', 'final_table_chestpain_k8_AHA.xlsx')
ap = argparse.ArgumentParser(); ap.add_argument('--out', default=TABLE); ap.add_argument('--check', action='store_true')
args = ap.parse_args()

# bundle -> triage (draft, from the bundle's entities and name, not from the literature)
TRIAGE = {'c0': 'Unclear', 'c1': 'Unclear', 'c2': 'Routine', 'c3': 'Manual RUW', 'c4': 'Routine', 'c5': 'Routine', 'c6': 'Routine', 'c7': 'Routine'}
# Excel row -> PMIDs kept (all rated Partly: no paper tests the order of the steps)
KEEP = {2: ['33320506', '15336583'], 3: ['15336583'], 4: ['33320506', '23418761', '31378324', '15336583'],
        5: ['34854348', '39235771', '26461998'], 6: ['34854348'], 7: ['34854348', '39235771', '26461998'],
        8: ['21534844'], 9: ['33964547'], 10: ['21534844'],
        11: ['18775237', '28925655', '36241466', '34709879'], 12: ['18775237', '28765984', '36241466', '34709879'],
        13: ['21534844', '28765984', '36241466', '34709879'],
        14: ['35123084', '15929749'], 15: ['30860925'], 16: ['30860925', '35123084'],
        17: ['24209723', '25111147', '27921113'], 18: ['24209723', '25111147', '27921113'], 19: ['24209723', '25111147', '27921113', '40860980'],
        20: ['30684966', '36205377'], 21: ['31864310', '30684966'], 22: ['15336583'],
        23: ['35735799'], 24: ['34854348'], 25: ['34854348']}
BASIS_TITLE = {'36241466', '34709879'}   # 2022 ACC ED consensus pathway (no abstract in PubMed) and the 2021 AHA/ACC guideline (abstract names no test)
STEP = {'CP': r'chest pain|chest discomfort|precordial', 'ecg': r'electrocardiogra|\bECGs?\b|\bEKGs?\b', 'trop': r'troponin',
        'cxr': r'chest radiograph|chest x-?ray|radiograph', 'CMP': r'metabolic panel', 'CBC': r'blood count|\bCBC\b',
        'venip': r'venipuncture|blood collection|blood draw|blood sampl', 'phys': r'physical examination|physical exam|bedside',
        'lipid': r'lipid|cholesterol', 'hyperl': r'hyperlipid|dyslipid|hypercholesterol', 'A1c': r'hba1c|glycated|glycosylated|hemoglobin a1c',
        'vitd': r'vitamin d|25-hydroxy', 'mpred': r'steroid|methylprednisolone|prednis', 'azithro': r'azithromycin',
        'gerd': r'reflux|GERD', 'endo': r'endoscop|EGD', 'path': r'biops|patholog|histolog',
        'manual': r'manual therapy|manipulat|physiotherap|physical therapy|chiropract', 'ther_ex': r'exercise|stretching',
        'pharyng': r'pharyngitis|sore throat|tonsillitis', 'cough': r'\bcough', 'uri': r'respiratory (tract )?infection|upper respiratory'}
abstracts = json.load(open(os.path.join(HERE, 'abstracts_AHA.json')))
returned = open(os.path.join(HERE, 'round1_AHA.log')).read() + open(os.path.join(HERE, 'round2_AHA.log')).read()

evidence, problems = [], []
for row, pmids in KEEP.items():
    steps = PAT[row]
    for pm in pmids:
        if f'{pm} ' not in returned:
            problems.append(f'row {row}: PMID {pm} was not returned by any search of this run')
        d = abstracts.get(pm)
        if not d:
            problems.append(f'row {row}: PMID {pm} has no fetched record'); continue
        text = d['title'] + ' ' + d['abstract']
        found = [t for t in dict.fromkeys(steps) if re.search(STEP[t], text, re.I)]
        basis = 'abstract' if len(found) >= 2 else ('title only' if pm in BASIS_TITLE else 'FAILS')
        if basis == 'FAILS':
            problems.append(f'row {row}: PMID {pm} names only {found} of {steps}')
        evidence.append([row, pm, d['year'], d['journal'], d['title'], ' '.join(steps), ' '.join(found),
                         ' '.join(t for t in dict.fromkeys(steps) if t not in found), basis])
    assert len(set(pmids)) == len(pmids)
if problems:
    raise SystemExit('\n'.join(problems))
print(f'{len(KEEP)} rows, {len({p for v in KEEP.values() for p in v})} distinct PMIDs, all returned in this run and all pass the step check')
if args.check:
    raise SystemExit(0)

wb = openpyxl.load_workbook(TABLE)
ws = wb.active
hdr = {c.value: c.column_letter for c in ws[1]}
colM, colO, colP = (hdr['Triage suggestion (Manual RUW / Routine / Unclear)'], hdr['Supported literature (Yes / Partly / None)'], hdr['PMIDs'])
assert (colM, colO, colP) == ('M', 'O', 'P')
assert all(ws[f'{c}{r}'].value is None for c in 'MOP' for r in range(2, 26)), 'the three columns are not empty: refusing to overwrite'
for r in range(2, 26):
    ws[f'{colO}{r}'] = 'Partly'
    ws[f'{colP}{r}'] = '; '.join(KEEP[r])
for r, b in zip(range(2, 26, 3), TRIAGE):
    assert ws[f'A{r}'].value.startswith(b + ' '), (r, ws[f'A{r}'].value[:20])
    ws[f'M{r}'] = TRIAGE[b]                        # top-left cell of the merged M{r}:M{r+2}
note = next(r for r in range(26, ws.max_row + 1) if str(ws[f'A{r}'].value).startswith('Triage suggestion, Supported literature'))
ws[f'A{note}'] = ("Triage suggestion (one per bundle, draft for underwriter review, from the bundle's entities and name, not from the literature): "
                  "Manual RUW = acute / ED work-up bundle; Routine = care that other, non-cardiac or unrelated care explains; Unclear = the bundle does not show the clinical context. "
                  "Supported literature and PMIDs (filled 2026-10-16): PubMed search of at least 3 of the bundle's top-5 entities ANDed (unordered) or the steps of the typical order of care ANDed, then the bundle's clinical setting; "
                  "a paper is kept only if its title fits the setting and its title or abstract names at least 2 of the row's steps (the 2021 AHA/ACC chest pain guideline and the 2022 ACC ED consensus pathway are kept on their titles). "
                  "PubMed cannot test order, so no row is Yes; Partly = the steps are addressed in a more general or different setting; every PMID was returned by a search of that run and checked in PubMed. "
                  "A clinician should confirm each Partly. Evidence per row: final_table_chestpain_k8_AHA_literature_evidence.csv.")
wb.save(args.out)
stem = args.out.rsplit('.', 1)[0]
with open(stem + '_literature_evidence.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['excel_row', 'PMID', 'year', 'journal', 'title (as returned by PubMed)', 'row steps', 'steps found in title/abstract', 'steps not found', 'basis'])
    w.writerows(evidence)
with open(stem + '_literature_search_log.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['round', 'kind', 'key', 'query', 'pubmed_hits'])
    for x in json.load(open(os.path.join(HERE, 'lit_results_AHA.json'))): w.writerow([1, x['kind'], x['key'] + ' ' + x['concepts'], x['query'], x['count']])
    for x in json.load(open(os.path.join(HERE, 'lit2_results_AHA.json'))): w.writerow([2, 'setting', x['label'], x['query'], x['count']])
print('wrote', args.out)
