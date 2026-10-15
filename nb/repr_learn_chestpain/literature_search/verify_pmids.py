"""Check every PMID in a final-table workbook against PubMed and write a verified-citation CSV.

usage: python verify_pmids.py final_table_letrozole_k7_full.xlsx
Reads the 'PMIDs' column (all rows), asks PubMed esummary for each ID, fails loudly if any ID does not exist,
and writes <workbook stem>_literature_verified.csv (PMID, first author, year, journal, title, table rows that use it).
"""
import csv, json, re, subprocess, sys
import openpyxl

path = sys.argv[1]
ws = openpyxl.load_workbook(path).active
col = next(c.column_letter for c in ws[1] if c.value == 'PMIDs')
rows_for = {}
for r in range(2, ws.max_row + 1):
    for p in re.findall(r'\b\d{6,8}\b', str(ws[f'{col}{r}'].value or '')):
        rows_for.setdefault(p, []).append(r)
ids = sorted(rows_for)
if not ids:
    sys.exit('no PMIDs in column ' + col)
url = 'https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?db=pubmed&retmode=json&id=' + ','.join(ids)
res = json.loads(subprocess.run(['/usr/bin/curl', '-sS', '-f', '-m', '60', url], capture_output=True, text=True, check=True).stdout)['result']
missing = [i for i in ids if i not in res or 'error' in res[i]]
if missing:
    sys.exit(f'PMIDs NOT FOUND in PubMed: {missing} (rows {[rows_for[i] for i in missing]})')
out = path.rsplit('.', 1)[0] + '_literature_verified.csv'
with open(out, 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['PMID', 'first_author', 'year', 'journal', 'title (as returned by PubMed)', 'used_in_table_rows'])
    for i in ids:
        d = res[i]
        w.writerow([i, d['authors'][0]['name'] if d.get('authors') else '', d['pubdate'][:4], d['source'], d['title'], ' '.join(map(str, rows_for[i]))])
print(f'{len(ids)} PMIDs verified -> {out}')
