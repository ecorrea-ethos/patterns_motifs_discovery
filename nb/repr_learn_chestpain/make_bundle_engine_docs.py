"""
Write one compact predicate-engine page per chest pain care pattern (NMF bundle) of the AHA k=8 run.

    python make_bundle_engine_docs.py            # needs pandas + openpyxl (env ethos_llm)

Inputs (nothing is measured here, every number is read from them; the first two are in outputs/):
  final_table_chestpain_k8_AHA.xlsx            top-5 entities, 3 ordered patterns per bundle, shares, verdicts, triage, PMIDs
  final_table_chestpain_k8_AHA_literature_verified.csv   title / author / year of each PMID in the table
  token_members_k8_chestpain_AHA.csv           concept ids behind every entity, size of the group tokens
Output: outputs/bundle_engines/Chestpain_c<N>_Predicate_Engine.md for c0, c1, c2, c4, c5, c6, c7.
c3 (ED work-up) is NOT written here: Chestpain_Workup_Predicate_Engine_v3.md is the page for it and defines the rules.

The rules are the v3 rules (window 1 day, 3 or more entities, one of 3 ordered sequences, chest pain required),
applied to the 5 most frequent entities and the 3 highest-lift patterns of each care pattern.
Hand-written part: ENT (wording of each entity), BUNDLES (which entities form one service, which patterns).
"""
import os, re
import pandas as pd
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUTPUTS = os.path.join(HERE, "outputs")
OUT = os.path.join(OUTPUTS, "bundle_engines")
os.makedirs(OUT, exist_ok=True)

ws = openpyxl.load_workbook(os.path.join(OUTPUTS, "final_table_chestpain_k8_AHA.xlsx")).active
members = pd.read_csv(os.path.join(HERE, "token_members_k8_chestpain_AHA.csv"), low_memory=False)
lit = pd.read_csv(os.path.join(OUTPUTS, "final_table_chestpain_k8_AHA_literature_verified.csv"), dtype=str)
N_COHORT = 139_416

# ------------------------------------------------------------------ the final table
TABLE = {}                                   # 'c0' -> dict(n, pct, triage, ents=[...], pats=[...])
cur = None
for r in range(2, ws.max_row + 1):
    g = ws.cell(r, 1).value
    if g and re.match(r"c\d ·", str(g)):
        cur = g.split(" ")[0]
        n, pct = re.match(r"([\d,]+) \(([\d.]+%)\)", str(ws.cell(r, 2).value)).groups()
        ents = []
        for blk in re.split(r"\n(?=\d\. )", str(ws.cell(r, 3).value)):
            head, _, tail = blk.partition("\n")
            name = re.sub(r"^\d\. ", "", head)
            share = int(re.search(r"\((\d+)% of bundle\)", name).group(1))
            name = re.sub(r" \(\d+% of bundle\)", "", name)
            if "same day as" in tail:
                med = 0
            else:
                d, side = re.search(r"median ([\d,]+) d (before|after)", tail).groups()
                med = int(d.replace(",", "")) * (-1 if side == "before" else 1)
            before = int(re.search(r"\((\d+)% of them before\)", tail).group(1))
            ents.append(dict(name=name, share=share, med=med, before=before))
        TABLE[cur] = dict(n=n, pct=pct, triage=ws.cell(r, 13).value, ents=ents, pats=[], lit_rating=ws.cell(r, 15).value)
    if cur and ws.cell(r, 5).value and ws.cell(r, 4).value:
        TABLE[cur]["pats"].append(dict(
            row=r, text=ws.cell(r, 5).value, sup_in=ws.cell(r, 7).value, sup_out=ws.cell(r, 8).value,
            passrate=ws.cell(r, 9).value, holds=ws.cell(r, 10).value, lift=ws.cell(r, 11).value,
            verdict=ws.cell(r, 12).value, order=ws.cell(r, 14).value, pmids=str(ws.cell(r, 16).value or "")))

# ------------------------------------------------------------------ wording of each entity (hand-written, ASD-STE style)
# key: (display name, [sentences], predicate name, event description, vocabulary, code or None, group token or None)
ENT = {
    "lipid": ("Lipid panel", ["A lipid panel is a blood test for cholesterol and other fats.", "Doctors order it at check-ups, so it alone does not show chest pain care."], "lipid_panel", "lipid panel", "CPT4", "80061", None),
    "venip": ("Venipuncture", ["Venipuncture is a blood draw from a vein.", "It does not show which test."], "venipuncture", "venipuncture for a blood test", "SNOMED", "28520004", None),
    "genexam": ("General examination", ["The general examination is the routine physical examination of a visit.", "Most patients at a check-up have it."], "general_exam", "general examination of the patient", "SNOMED", "162673000", None),
    "ecg": ("ECG", ["The ECG (electrocardiogram) is a heart tracing.", "The list holds 5 ECG codes, with or without the report."], "ecg", "ECG (any of 5 codes)", None, None, "ekg"),
    "cmp": ("CMP", ["The CMP (comprehensive metabolic panel) is a blood chemistry test.", "It measures kidney function, liver function, and blood salts."], "cmp", "comprehensive metabolic panel", "SNOMED", "371361000119107", None),
    "cbc": ("CBC", ["The CBC is the complete blood count with a white cell count.", "It counts the cells of the blood."], "cbc", "complete blood count with white cell count", "SNOMED", "9564003", None),
    "hld": ("Hyperlipidemia", ["Hyperlipidemia is a high level of fats in the blood.", "The list holds 6 diagnosis codes."], "hyperlipidemia", "hyperlipidemia diagnosis (any of 6 codes)", None, None, "hyperlipidemia"),
    "hba1c": ("HbA1c", ["HbA1c is a blood test for the average blood sugar of the last 3 months.", "The list holds 1 code."], "hba1c", "HbA1c test", None, None, "hba1c"),
    "methylpred": ("Methylprednisolone", ["Methylprednisolone is a steroid drug.", "The code is the 21-tablet dose pack."], "methylprednisolone", "methylprednisolone 21-tablet pack", "RxNorm", "762675", None),
    "amoxclav": ("Amoxicillin with clavulanate", ["Amoxicillin with clavulanate is an antibiotic.", "The code is the tablet of 875 mg with 125 mg."], "amoxicillin_clavulanate", "amoxicillin 875 mg with clavulanate 125 mg tablet", "RxNorm", "562508", None),
    "azithro": ("Azithromycin", ["Azithromycin is an antibiotic.", "The code is the pack of 6 tablets of 250 mg."], "azithromycin", "azithromycin 6-tablet pack", "RxNorm", "749783", None),
    "ppi": ("PPI or H2 blocker", ["A PPI or an H2 blocker is a drug that lowers stomach acid.", "The list holds 85 drug codes."], "ppi_h2", "PPI or H2 blocker drug (any of 85 codes)", None, None, "ppi_h2"),
    "pred": ("Prednisone", ["Prednisone is a steroid drug.", "The code is the tablet of 20 mg."], "prednisone", "prednisone 20 mg tablet", "RxNorm", "312615", None),
    "gyn": ("Gynecologic examination", ["The gynecologic examination is the pelvic examination of a woman.", "It is part of women's check-ups."], "gyn_exam", "gynecologic examination", "SNOMED", "83607001", None),
    "sti": ("STI test", ["The STI test finds the DNA or RNA of an infection that passes by sex.", "The list holds the chlamydia code and the gonorrhea code. The engine counts them as one service."], "sti_nucleic_acid_test", "STI nucleic acid test (chlamydia OR gonorrhea)", None, None, None),
    "cyto": ("Cervical cytology", ["Cervical cytology is a test of cells from the cervix or the vagina.", "It is also the Pap test."], "cervical_cytology", "cervical or vaginal cytology", "CPT4", "88175", None),
    "upt": ("Urine pregnancy test", ["The urine pregnancy test is a quick test for a pregnancy."], "urine_pregnancy_test", "urine pregnancy test", "SNOMED", "104405006", None),
    "path": ("Surgical pathology", ["Surgical pathology is the study of a tissue sample by eye and under a microscope.", "A biopsy gives such a sample."], "surgical_pathology", "surgical pathology level IV", "CPT4", "88305", None),
    "crc": ("Colon cancer screening", ["Colon cancer screening is a test for colon cancer in a person with no symptoms."], "colon_cancer_screening", "screening for malignant neoplasm of colon", "SNOMED", "275978004", None),
    "egd": ("Upper endoscopy with biopsy", ["Upper endoscopy is a look into the food pipe and the stomach with a camera.", "The doctor also takes a tissue sample."], "upper_endoscopy_biopsy", "endoscopy and biopsy of the upper gastrointestinal tract", "SNOMED", "445982009", None),
    "gerd": ("GERD", ["GERD is acid reflux from the stomach into the food pipe.", "The list holds 4 diagnosis codes, and heartburn is one of them."], "gerd", "GERD diagnosis (any of 4 codes)", None, None, "gerd"),
    "periodic": ("Periodic preventive visit", ["The periodic preventive visit is the check-up of an adult of 40 to 64 years."], "periodic_preventive_visit", "periodic comprehensive preventive visit (99396)", "CPT4", "99396", None),
    "therproc": ("Therapeutic procedure", ["A therapeutic procedure is physical therapy in units of 15 minutes.", "The list holds the exercise code (97110) and the movement training code (97112). The engine counts them as one service."], "therapeutic_procedure", "therapeutic procedure (97110 OR 97112)", None, None, None),
    "trop": ("Troponin", ["Troponin is a blood test for injury to the heart muscle.", "The data does not show the type of test. A high-sensitivity test and a standard test look the same."], "troponin", "troponin test", None, None, "troponin"),
    "lipase": ("Lipase", ["Lipase is a blood test for the pancreas.", "Doctors also use it for abdominal pain."], "lipase", "lipase test", "SNOMED", "104988000", None),
    "cxr": ("Chest X-ray", ["The list holds X-rays of 1, 2, 3, and 4 or more views."], "cxr", "chest X-ray", None, None, "chest_xray"),
    "ivpush": ("IV push", ["An IV push puts a drug into a vein with a syringe.", "The list holds the first drug and each additional drug. The engine counts them as one service."], "iv_push", "IV push (first drug OR each additional drug)", None, None, None),
    "ther97110": ("Therapeutic exercise", ["Therapeutic exercise is physical therapy with exercises, in units of 15 minutes.", "It is one code of the therapeutic procedure list."], "therapeutic_exercise", "therapeutic exercise (97110)", "CPT4", "97110", None),
    "manual": ("Manual therapy", ["Manual therapy is hands-on treatment by a therapist, for example joint movement."], "manual_therapy", "manual therapy techniques", "CPT4", "97140", None),
    "theract": ("Therapeutic activities", ["Therapeutic activities are one-on-one physical therapy with movements of daily life."], "therapeutic_activities", "therapeutic activities", "CPT4", "97530", None),
    "lbp": ("Low back pain", ["Low back pain is a diagnosis of pain in the lower back."], "low_back_pain", "low back pain diagnosis", "SNOMED", "279039007", None),
    "b12": ("Vitamin B12", ["Vitamin B12 is a blood test for the level of this vitamin."], "vitamin_b12", "vitamin B12 measurement", "SNOMED", "14598005", None),
    "vitd": ("Vitamin D", ["The vitamin D test measures 25-hydroxy vitamin D in the blood."], "vitamin_d", "25-hydroxy vitamin D measurement", "SNOMED", "12199005", None),
    "t4": ("Free T4", ["Free T4 is a blood test for the level of the thyroid hormone."], "free_t4", "free T4 measurement", "SNOMED", "5113004", None),
    "iron": ("Iron", ["The iron test measures the iron of the blood."], "iron", "iron measurement", "SNOMED", "42950004", None),
    "ferritin": ("Ferritin", ["Ferritin is a blood test for the iron stores of the body."], "ferritin", "ferritin measurement", "SNOMED", "489004", None),
    "cough": ("Cough", ["Cough is a diagnosis of a cough."], "cough", "cough diagnosis", "SNOMED", "49727002", None),
    "uri": ("Acute upper respiratory infection", ["Acute upper respiratory infection is a cold or a similar infection of the nose and throat."], "upper_respiratory_infection", "acute upper respiratory infection diagnosis", "SNOMED", "54398005", None),
    "pharyngitis": ("Acute pharyngitis", ["Acute pharyngitis is a sore throat from an infection."], "pharyngitis", "acute pharyngitis diagnosis", "SNOMED", "363746003", None),
    "antigen": ("Rapid antigen test", ["A rapid antigen test finds an infection by a protein of the germ.", "The list holds 2 codes (87880 and 87804). The engine counts them as one service."], "rapid_antigen_test", "rapid antigen test (87880 OR 87804)", None, None, None),
}
# codes that stand behind an entity that is a union of several table entities
UNION_CODES = {"ivpush": [("CPT4", "96374"), ("CPT4", "96375")], "sti": [("CPT4", "87491"), ("CPT4", "87591")], "therproc": [("CPT4", "97110"), ("CPT4", "97112")],
               "antigen": [("CPT4", "87880"), ("CPT4", "87804")]}


# noun phrase of each entity: "The applicant has at least one <noun>." and the predicate definition
NOUN = {"trop": "troponin test", "lipase": "lipase test", "cxr": "chest X-ray", "ivpush": "IV push", "ther97110": "therapeutic exercise session", "lipid": "lipid panel", "venip": "venipuncture", "genexam": "general examination", "ecg": "ECG", "cmp": "CMP test", "cbc": "CBC test",
        "hld": "hyperlipidemia diagnosis", "hba1c": "HbA1c test", "methylpred": "methylprednisolone prescription",
        "amoxclav": "amoxicillin with clavulanate prescription", "azithro": "azithromycin prescription", "ppi": "PPI or H2 blocker prescription",
        "pred": "prednisone prescription", "gyn": "gynecologic examination", "sti": "STI test", "cyto": "cervical cytology test",
        "upt": "urine pregnancy test", "path": "surgical pathology exam", "crc": "colon cancer screening", "egd": "upper endoscopy with biopsy",
        "gerd": "GERD diagnosis", "periodic": "periodic preventive visit", "therproc": "therapeutic procedure", "manual": "manual therapy session",
        "theract": "therapeutic activities session", "lbp": "low back pain diagnosis", "b12": "vitamin B12 test", "vitd": "vitamin D test",
        "t4": "free T4 test", "iron": "iron test", "ferritin": "ferritin test", "cough": "cough diagnosis",
        "uri": "acute upper respiratory infection diagnosis", "pharyngitis": "acute pharyngitis diagnosis", "antigen": "rapid antigen test"}
AN_FIRST = set("AEFHILMNORSX")      # letters whose name starts with a vowel sound, for acronyms


def with_article(noun):
    w = noun.split()[0]
    if w.isupper() or any(c.isupper() for c in w[1:]):         # acronym: ECG, CMP, HbA1c
        return ("an " if w[0] in AN_FIRST else "a ") + noun
    noun = noun[0].lower() + noun[1:]
    return ("an " if noun[0] in "aeiou" and not noun.startswith("ur") else "a ") + noun


def phrase(steps, j):
    k = steps[j]
    if k == "cp":
        return "a chest pain event"
    if k == "x2":
        return "a second " + NOUN[steps[j - 1]].replace("ECG", "ECG")
    return with_article(NOUN[k])


def lower_first(t):
    w = t.split()[0]
    return t if (w.isupper() or any(c.isupper() for c in w[1:])) else t[0].lower() + t[1:]

# ------------------------------------------------------------------ per care pattern: services (entities that count) and the 3 patterns
# services: list of (entity key, indices of the final-table entities that it stands for)
# pats: the 3 table patterns, as step lists ('cp' = chest pain; 'x2' repeats the previous step as a second event)
BUNDLES = {
    "c3": dict(prefix="ed_workup", name="ED work-up",
               what="This care pattern is emergency department (ED) care for acute chest pain. Troponin, chest X-ray, lipase, and IV drugs dominate it.",
               services=[("trop", [0]), ("ivpush", [1, 4]), ("lipase", [2]), ("cxr", [3])],
               pats=[["cp", "cxr", "trop", "cmp"], ["cp", "ecg", "trop", "cmp"], ["cp", "trop", "cmp", "cbc"]]),
    "c0": dict(prefix="checkup", name="outpatient check-up and elective cardiac testing",
               what="This care pattern is routine outpatient care. Lipid tests, ECG, and check-up visits dominate it.",
               services=[("lipid", [0]), ("venip", [1]), ("genexam", [2]), ("ecg", [3]), ("cmp", [4])],
               pats=[["cp", "ecg", "hld"], ["cp", "genexam", "lipid"], ["cp", "ecg", "genexam"]]),
    "c1": dict(prefix="rx", name="prescriptions for acute and chronic illness",
               what="This care pattern is drug care for short illness and long illness. Steroids, antibiotics, and acid drugs dominate it.",
               services=[("methylpred", [0]), ("amoxclav", [1]), ("azithro", [2]), ("ppi", [3]), ("pred", [4])],
               pats=[["methylpred", "cp", "ecg"], ["azithro", "cp", "ecg"], ["cp", "ecg", "methylpred"]]),
    "c2": dict(prefix="womens", name="women's health and urinary tract",
               what="This care pattern is women's health care and urinary care. Pelvic examination and infection tests dominate it.",
               services=[("gyn", [0]), ("sti", [1, 2]), ("cyto", [3]), ("upt", [4])],
               pats=[["cp", "cbc", "venip"], ["cp", "venip", "x2"], ["venip", "cp", "ecg"]]),
    "c4": dict(prefix="gi", name="stomach and bowel work-up and cancer screening",
               what="This care pattern is care of the stomach and the bowel. Biopsy, pathology, and screening dominate it.",
               services=[("path", [0]), ("crc", [1]), ("egd", [2]), ("gerd", [3]), ("periodic", [4])],
               pats=[["cp", "gerd", "egd"], ["cp", "path", "egd"], ["cp", "gerd", "path"]]),
    "c5": dict(prefix="msk", name="muscle and bone care and physical therapy",
               what="This care pattern is physical therapy for pain of muscle and bone. Therapy codes and back pain dominate it.",
               services=[("therproc", [0, 4]), ("manual", [1]), ("theract", [2]), ("lbp", [3])],
               pats=[["cp", "manual", "ther97110"], ["cp", "manual", "x2"], ["cp", "ther97110", "manual"]]),
    "c6": dict(prefix="labs", name="nutrition and hormone blood tests",
               what="This care pattern is blood testing for vitamins, iron, and hormones. Lipid and HbA1c tests follow chest pain in its sequences.",
               services=[("b12", [0]), ("vitd", [1]), ("t4", [2]), ("iron", [3]), ("ferritin", [4])],
               pats=[["cp", "lipid", "vitd"], ["cp", "lipid", "hba1c"], ["cp", "ecg", "lipid"]]),
    "c7": dict(prefix="resp", name="acute respiratory infection",
               what="This care pattern is care of a cold, a sore throat, or a similar infection. Cough, infection diagnoses, and rapid tests dominate it.",
               services=[("cough", [0]), ("uri", [1]), ("pharyngitis", [2]), ("antigen", [3, 4])],
               pats=[["pharyngitis", "cp", "ecg"], ["cp", "ecg", "uri"], ["cough", "cp", "ecg"]]),
}
STEP_ENT = {"cp": None}


def cid_of(vocab, code):
    s = members[(members.vocabulary_id == vocab) & (members.concept_code.astype(str) == code)]
    assert len(s) >= 1, (vocab, code)
    return int(s.concept_id.iloc[0]), str(s.concept_name.iloc[0])


def group_info(tok):
    s = members[members.token == tok]
    assert len(s) >= 1, tok
    return int(len(s))


def code_text(k):
    """Short code description of an entity for the predicate table."""
    name, sents, pred, desc, vocab, code, tok = ENT[k]
    if k in UNION_CODES:
        parts = [f"{v} {c} (concept_id {cid_of(v, c)[0]})" for v, c in UNION_CODES[k]]
        return " + ".join(parts)
    if tok:
        return f"group `{tok}`, {group_info(tok)} concepts (see the member table)"
    return f"{vocab} {code} (concept_id {cid_of(vocab, code)[0]})"


def seq_name(steps):
    keep = steps if steps[0] != "cp" else steps[1:]
    return "composite_seq_" + "_".join(keep)


def step_pred(k, steps, i):
    if k == "cp":
        return "chest_pain"
    if k == "x2":
        return ENT[steps[i - 1]][2]
    return ENT[k][2]


def step_label(steps, i):
    k = steps[i]
    if k == "cp":
        return "chest pain"
    if k == "x2":
        return "a second " + ENT[steps[i - 1]][0] + " event"
    return ENT[k][0]


ORD = ["first", "second", "third"]


def fmt_pct(x):
    return f"{float(x) * 100:.1f}%"


def days_text(med):
    return "same day" if med == 0 else f"{abs(med):,} d {'before' if med < 0 else 'after'}"


def ent_labels(cn):
    """Readable label + code for each of the 5 table entities."""
    out = []
    for i, e in enumerate(TABLE[cn]["ents"]):
        k = next(k for k, idx in BUNDLES[cn]["services"] if i in idx)
        m = re.search(r"\[(.*?)\]", e["name"])
        code = (m.group(1).replace("code group", "code group " + ENT[k][6]) if ENT[k][6] else m.group(1)) if m else ""
        out.append(f"{ENT[k][0]} ({code})" if code else ENT[k][0])
    return out


def doc(cn):
    B, T = BUNDLES[cn], TABLE[cn]
    p = B["prefix"]
    n_ent = len(B["services"])
    L = []
    a = L.append
    a(f"# Chest Pain Care Pattern {cn}: {B['name'][0].upper() + B['name'][1:]}")
    a("")
    a("*Working Draft — October 2026*")
    a("")
    a("---")
    a("")
    a("## Overview")
    a("")
    a(f"This document describes the predicates and composites for care pattern {cn} of the chest pain predicate engine. "
      f"The cohort has {N_COHORT:,} chest pain patients. A data method groups these patients by the care they receive. "
      f"The method found 8 care patterns. Care pattern {cn} is **{B['name']}**. It has {T['n']} patients ({T['pct']}).")
    a("")
    a(f"{B['what']} The ED work-up engine ([Chestpain_Workup_Predicate_Engine_v3.md](../Chestpain_Workup_Predicate_Engine_v3.md)) defines the rules. "
      f"This document applies the same rules to the 5 most frequent entities and to the 3 most specific ordered patterns of care pattern {cn}. "
      "It contains no underwriting logic. It is a draft and has no validation yet.")
    a("")
    meds = [e["med"] for e in T["ents"]]
    lo, hi = min(meds), max(meds)
    fit = (f"The first event of each of the 5 entities has a median between {days_text(lo)} and {days_text(hi)} against the first chest pain. "
           if lo != hi else "")
    far = sum(1 for m in meds if abs(m) > 1)
    used = {q for steps in B["pats"] for q in steps}
    overlap = any(k in used for k, _ in B["services"])
    far_txt = ("All 5 entities" if far == 5 else f"{far} of the 5 entities") + " have a median more than 1 day from the first chest pain. " if far else ""
    a(f"The draft triage for this care pattern is **{T['triage']}**. It comes from the entities and the name of the care pattern. It does not come from the literature. "
      f"An underwriter reviews it.")
    a("")
    a(f"{fit}{far_txt}" + ("The window of 1 day is short against these medians. The share of patients in the window is not measured. The window needs a calibration before any use." if far else "The window of 1 day fits the timing of the 5 entities."))
    if not overlap:
        a("")
        a("The steps of the 3 sequences are not among the 5 entities. They are general steps of care.")
    a("")
    a("---")
    a("")
    a("## Predicates and Composites")
    a("")
    a("Each predicate uses a list of codes. The code of each list is in the text. The member table `token_members_k8_chestpain_AHA.csv` shows every concept of a code group.")
    a("")
    a("### Atomic Predicates")
    a("")
    a("A predicate answers one question about the event history of an applicant. It returns true or false. The code and the date are attributes of one event.")
    a("")
    a("**Chest pain present.** The applicant has at least one chest pain event. The list holds chest pain (77670), precordial pain (134159), chest pain on breathing (4168213), and pain of the space between the ribs (42539051).")
    a("")
    a("`chest_pain_present = event IS chest pain diagnosis`")
    a("")
    a("**First chest pain date.** This is the date of the first chest pain event. The engine measures the other dates from this date.")
    a("")
    a("`first_chest_pain_date = MIN(event_date) over chest pain events`")
    a("")
    shown = set()
    for k, idx in B["services"]:
        name, sents, pred, desc, vocab, code, tok = ENT[k]
        a(f"**{name} present.** The applicant has at least one {NOUN[k]}. {' '.join(sents)} Code: {code_text(k)}.")
        a("")
        a(f"`{pred}_present = event IS {desc}`")
        a("")
        shown.add(k)
    extra = []
    for steps in B["pats"]:
        for s in steps:
            if s not in ("cp", "x2") and s not in shown and s not in extra:
                extra.append(s)
    for k in extra:
        name, sents, pred, desc, vocab, code, tok = ENT[k]
        a(f"**{name} present.** The applicant has at least one {NOUN[k]}. {' '.join(sents)} It is a step of a sequence, not one of the {n_ent} entities. Code: {code_text(k)}.")
        a("")
        a(f"`{pred}_present = event IS {desc}`")
        a("")
    a(f"**Entity in window.** An entity is in the window when its event is not more than 1 day before or after a chest pain event. The window is 1 day. It is the window of the ED work-up engine.")
    a("")
    a("`workup_window_days = 1`")
    a("")
    a("`<entity>_in_window = <entity>_present AND EXISTS chest pain event date d, entity event date e with |e - d| <= workup_window_days`")
    a("")
    a(f"There is one predicate for each of the {n_ent} entities: " + ", ".join(f"`{ENT[k][2]}_in_window`" for k, _ in B["services"]) + ".")
    a("")
    a("---")
    a("")
    a("### Composites")
    a("")
    a("A composite combines atomic predicates, or other composites, into a named clinical pattern. Composites are in three groups: entities, sequences, and the care pattern.")
    a("")
    a("#### Entities")
    a("")
    a(f"**Entity count.** This is the number of the {n_ent} entities that are in the window for the applicant.")
    a("")
    a(f"`{p}_entity_count = " + " + ".join(f"{ENT[k][2]}_in_window" for k, _ in B["services"]) + "`")
    a("")
    a(f"**Entities.** The applicant has 3 or more of the {n_ent} entities in the window. The value 3 is a draft. It has no calibration yet.")
    a("")
    a(f"`composite_{p}_entities = {p}_entity_count >= 3`")
    a("")
    a("#### Sequences")
    a("")
    all_cp_first = all(s[0] == "cp" for s in B["pats"])
    a("A sequence is a list of events in a fixed order. Other events between the steps do not break the sequence. "
      + ("Every sequence starts at a chest pain event. The last step is not more than 1 day after that chest pain event. "
         if all_cp_first else
         "Every sequence has one chest pain step. A step after the chest pain step is not more than 1 day after that chest pain event. A step before the chest pain step has no time limit. ")
      + "A step is true for any code of its list.")
    a("")
    a("`sequence_window_days = 1`")
    a("")
    a(f"The data method found these 3 sequences. They are the 3 patterns of care pattern {cn} with the highest lift. Each sequence has 3 steps."
      + (""))
    a("")
    names = []
    for i, steps in enumerate(B["pats"]):
        nm = seq_name(steps)
        names.append(nm)
        ph = [phrase(steps, j) for j in range(3)]
        a(f"**Sequence {'ABC'[i]}.** The applicant has {ph[0]}, then {ph[1]}, then {ph[2]}.")
        a("")
        a("`" + nm + " = " + " -> ".join(step_pred(s_, steps, j) for j, s_ in enumerate(steps)) + "`")
        a("")
    a(f"**Sequence of care pattern {cn}.** The applicant has at least one of the sequences A, B, or C.")
    a("")
    a(f"`composite_{p}_sequence = " + " OR ".join(names) + "`")
    a("")
    a("#### Same-day order")
    a("")
    a("Claims have a date and no time. The engine puts the events of one day in a fixed order: diagnosis, procedure, lab test, then drug. "
      "Inside one type, a lower code ID comes first. For a same-day pattern, the order of the steps is a rule. It is not the clock time.")
    a("")
    a("#### Care pattern")
    a("")
    a(f"**Care pattern {cn}.** The applicant has chest pain, 3 or more entities in the window, and at least one sequence. "
      "This pattern shows the type of care. It does not show a diagnosis. The predicates test that an event happened, not the result.")
    a("")
    a(f"`composite_{p}_pattern = chest_pain_present AND composite_{p}_entities AND composite_{p}_sequence`")
    a("")
    a("---")
    a("")
    a("## Population Coverage")
    a("")
    a(f"Counts come from the chest pain cohort of {N_COHORT:,} patients. The tables come from the k=8 run of the data method (`nmf_k8_chestpain_AHA`). "
      "The engine has no measured rate for any predicate or composite yet.")
    a("")
    a("### Entity rates")
    a("")
    a(f"The rates are for the {T['n']} patients of care pattern {cn}. They count an entity at any time, not only in the window.")
    a("")
    a("| Entity | Patients with the entity | Median day against first chest pain | Share that had it before first chest pain |")
    a("| :---- | ----: | :---- | ----: |")
    for e, lab in zip(T["ents"], ent_labels(cn)):
        a(f"| {lab} | {e['share']}% | {days_text(e['med'])} | {e['before']}% |")
    a("")
    a("### Sequence rates")
    a("")
    a("The shares count the steps in this order at any time. They do not use the 1-day limit. The other care patterns are a random sample of 1,500 patients.")
    a("")
    a(f"| Sequence | Patients of {cn} | Patients of the other care patterns | Times as common | Held-out pass rate | Verdict |")
    a("| :---- | ----: | ----: | ----: | ----: | :---- |")
    for i, pt in enumerate(T["pats"]):
        steps = B["pats"][i]
        parts = ["chest pain" if q == "cp" else ("second " + NOUN[steps[j - 1]]) if q == "x2" else NOUN[q] for j, q in enumerate(steps)]
        txt = ", ".join(parts)
        txt = txt[0].upper() + txt[1:] if not txt.split()[0].isupper() and not any(c.isupper() for c in txt.split()[0][1:]) else txt
        order = " (same-day order)" if pt["order"] == "Same-day bundle" else ""
        a(f"| {'ABC'[i]}. {txt}{order} | {fmt_pct(pt['sup_in'])} | {fmt_pct(pt['sup_out'])} | {float(pt['lift']):.1f} | {float(pt['passrate']):.0%} | {pt['verdict']} |")
    a("")
    n_hold = sum(1 for pt in T["pats"] if pt["holds"] == "Yes")
    if n_hold == 0:
        a("The held-out check repeats the test on random halves of the patients, in 40 runs. No sequence of this care pattern passed it. Read the sequences as a draft.")
    elif n_hold == 3:
        a("The held-out check repeats the test on random halves of the patients, in 40 runs. All 3 sequences passed it.")
    else:
        a(f"The held-out check repeats the test on random halves of the patients, in 40 runs. {n_hold} of the 3 sequences passed it.")
    a("")
    a("---")
    a("")
    a("## References")
    a("")
    a("- Gulati M, et al. *2021 AHA/ACC/ASE/CHEST/SAEM/SCCT/SCMR Guideline for the Evaluation and Diagnosis of Chest Pain.* Circulation 2021;144:e368. doi:10.1161/CIR.0000000000001029.")
    pm = []
    for pt in T["pats"]:
        for x in re.split(r"[;,\s]+", pt["pmids"]):
            if x.isdigit() and x not in pm:
                pm.append(x)
    for x in pm:
        r = lit[lit.PMID == x]
        if len(r):
            r = r.iloc[0]
            a(f"- {r['first_author']} et al. *{str(r['title (as returned by PubMed)']).rstrip('.')}.* {r['journal']} {r['year']}. PMID {x}.")
    if pm:
        a("")
        a(f"The papers above come from the literature check of the final table (rating: {T['lit_rating']}). They cover the care type. They do not test the order of the steps.")
    return "\n".join(L) + "\n"


if __name__ == "__main__":
    for cn in [c for c in BUNDLES if c != "c3"]:        # c3 has its own page: Chestpain_Workup_Predicate_Engine_v3.md
        path = os.path.join(OUT, f"Chestpain_{cn}_Predicate_Engine.md")
        open(path, "w").write(doc(cn))
        print("wrote", path)
