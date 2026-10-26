"""Fill Triage suggestion (M), Supported literature (O) and PMIDs (P) of outputs/final_table_colonoscopy_k7_v1.xlsx and write the audit files.

usage (conda env ethos_llm, cwd = literature_search/):
    python fill_literature.py --check              # test the decisions against the abstracts, write nothing
    python fill_literature.py [--out PATH.xlsx]    # write the workbook (default: the table itself, in place)

The table lists the NMF bundles c0-c6 (21 rows, Excel rows 2-22). PAT in pubmed.py holds the steps of each row; if cell 48 of the
notebook is re-run and the patterns change, update PAT, re-run round1.py / round2.py, fetch_abstracts.py, then KEEP below.
Inputs: round1.log / round2.log (what PubMed returned in this run), abstracts.json (efetch of every kept paper).
A row keeps a paper only if its title or abstract names at least min(2, number of distinct steps) of the row's steps (steps.py) and the
title fits the bundle's setting. PubMed cannot test order, so no row is rated Yes. A row with no kept paper is None / 'none found'.
Writes next to the workbook: *_literature_evidence.csv (per row and paper: steps found / not found), *_literature_search_log.csv.
"""
import argparse, csv, json, os, re
import openpyxl
from pubmed import PAT
from steps import STEP

HERE = os.path.dirname(os.path.abspath(__file__))
TABLE = os.path.join(HERE, '..', 'outputs', 'final_table_colonoscopy_k7_v1.xlsx')
FIRST, LAST = 2, 22
ap = argparse.ArgumentParser(); ap.add_argument('--out', default=TABLE); ap.add_argument('--check', action='store_true')
args = ap.parse_args()

# bundle (first row of its 3-row block) -> triage (draft, from the bundle's entities and name, not from the literature)
TRIAGE = {2: 'Manual RUW', 5: 'Routine', 8: 'Routine', 11: 'Routine', 14: 'Routine', 17: 'Unclear', 20: 'Routine'}
# table row -> PMIDs kept (all Partly: no paper tests the order of the steps). Rows missing here have no kept paper -> None.
# None rows: 2-4 (c0: the 5-fluorouracil / capecitabine / oxaliplatin cardiotoxicity papers name the ECG but no chemistry or blood-count step),
#            8-10 (c2: no paper names a PPI / H2 blocker together with the antibiotic), 13 (c3: no paper names the preventive visit and the pelvic exam).
KEEP = {5: ['34100866', '30586131', '28671929'], 6: ['30586131', '26430473', '11277547'],
        7: ['34100866', '23922061', '30488532', '31100143'],
        11: ['40489104', '35469700', '32059094', '36033808'], 12: ['40489104', '35469700', '32059094'],
        14: ['29373057', '34152897', '32723399', '35881705'], 15: ['29373057', '34152897', '32723399', '35881705'],
        16: ['29373057', '34152897', '32723399', '35881705'],
        17: ['18416471', '38670139', '32572858', '37737817'], 18: ['38670139', '28212588', '32572858', '37737817'],
        19: ['18416471', '25278696', '32572858', '37737817'],
        20: ['40895625', '35743024', '39107936', '30323637'], 21: ['35801218', '40895625', '35743024', '35498428'],
        22: ['30323637', '35801218', '39107936']}
abstracts = json.load(open(os.path.join(HERE, 'abstracts.json'))) if os.path.exists(os.path.join(HERE, 'abstracts.json')) else {}
returned = open(os.path.join(HERE, 'round1.log')).read() + open(os.path.join(HERE, 'round2.log')).read()

evidence, problems = [], []
for row, pmids in KEEP.items():
    steps = list(dict.fromkeys(PAT[row])); need = min(2, len(steps))
    assert len(set(pmids)) == len(pmids)
    for pm in pmids:
        if f'{pm} ' not in returned:
            problems.append(f'row {row}: PMID {pm} was not returned by any search of this run')
        d = abstracts.get(pm)
        if not d:
            problems.append(f'row {row}: PMID {pm} has no fetched record'); continue
        text = d['title'] + ' ' + d['abstract']
        found = [t for t in steps if re.search(STEP[t], text, re.I)]
        basis = ('abstract' if d['abstract'] else 'title only') if len(found) >= need else 'FAILS'
        if basis == 'FAILS':
            problems.append(f'row {row}: PMID {pm} names only {found} of {steps}')
        evidence.append([row, pm, d['year'], d['journal'], d['title'], ' '.join(steps), ' '.join(found), ' '.join(t for t in steps if t not in found), basis])
if problems:
    raise SystemExit('\n'.join(problems))
n_none = [r for r in range(FIRST, LAST + 1) if r not in KEEP]
print(f'{len(KEEP)} rows Partly, {len(n_none)} rows None {n_none}; {len({p for v in KEEP.values() for p in v})} distinct PMIDs, all returned in this run and all pass the step check')
if args.check:
    raise SystemExit(0)

wb = openpyxl.load_workbook(TABLE)
ws = wb.active
hdr = {c.value: c.column_letter for c in ws[1]}
colM, colO, colP = hdr['Triage suggestion (Manual RUW / Routine / Unclear)'], hdr['Supported literature (Yes / Partly / None)'], hdr['PMIDs']
assert (colM, colO, colP) == ('M', 'O', 'P')
assert all(ws[f'{c}{r}'].value is None for c in 'MOP' for r in range(FIRST, LAST + 1)), 'the three columns are not empty: refusing to overwrite'
assert str(ws[f'A{FIRST}'].value).startswith('c0 ') and str(ws[f'A{LAST - 2}'].value).startswith('c6 '), 'the table must list only the bundles c0-c6 (re-run cell 48 of the notebook)'
assert sorted(TRIAGE) == list(range(FIRST, LAST, 3)), 'one triage per bundle'
for r in range(FIRST, LAST + 1):
    ws[f'O{r}'] = 'Partly' if r in KEEP else 'None'
    ws[f'P{r}'] = '; '.join(KEEP[r]) if r in KEEP else 'none found'
for r, t in TRIAGE.items():
    assert any(str(m) == f'M{r}:M{r + 2}' for m in ws.merged_cells.ranges), r
    ws[f'M{r}'] = t                                   # top-left cell of the merged M{r}:M{r+2}
note = next(r for r in range(LAST + 1, ws.max_row + 1) if str(ws[f'A{r}'].value).startswith('Triage suggestion, Supported literature'))
ws[f'A{note}'] = ("Triage suggestion (one per bundle, draft for underwriter review, from the bundle's entities and name, not from the literature): "
                  "Manual RUW = colorectal cancer treatment shown (c0: ostomy, colectomy, chemotherapy and radiation load on it and it is enriched for cancer, OR 2.7); "
                  "Routine = care that screening, check-ups or unrelated care explains (c1, c2, c3, c4, c6); "
                  "Unclear = polypectomy episode where the codes do not show the polyp size, number or histology that the colonoscopy underwriting note asks for (c5). "
                  "Supported literature and PMIDs (filled 2026-10-23): PubMed search of at least 3 of the bundle's top-5 entities ANDed (unordered) or the steps of the typical order of care ANDed, "
                  "then the bundle's clinical setting (acute care and colorectal cancer treatment, cancer screening, bowel preparation, polypectomy surveillance, rehabilitation, iron deficiency); "
                  "a paper is kept only if its title fits the setting and its title or abstract names at least 2 of the row's steps "
                  "(1 when the row has a single distinct step; the 'chemistry panel' step also counts liver function tests, the 'venipuncture' step also counts blood tests). "
                  "PubMed cannot test order, so no row is Yes; Partly = the steps are addressed in a more general or different setting; "
                  "None = no paper named the steps ('none found'); every PMID was returned by a search of that run and checked in PubMed. A clinician should confirm each Partly. "
                  "Evidence per row: final_table_colonoscopy_k7_v1_literature_evidence.csv.")
wb.save(args.out)
stem = args.out.rsplit('.', 1)[0]
with open(stem + '_literature_evidence.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['excel_row', 'PMID', 'year', 'journal', 'title (as returned by PubMed)', 'row steps', 'steps found in title/abstract', 'steps not found', 'basis'])
    w.writerows(evidence)
with open(stem + '_literature_search_log.csv', 'w', newline='') as f:
    w = csv.writer(f); w.writerow(['round', 'kind', 'key', 'query', 'pubmed_hits'])
    for x in json.load(open(os.path.join(HERE, 'lit_results.json'))): w.writerow([1, x['kind'], x['key'] + ' ' + x['concepts'], x['query'], x['count']])
    for x in json.load(open(os.path.join(HERE, 'lit2_results.json'))): w.writerow([2, 'setting', x['label'], x['query'], x['count']])
print('wrote', args.out)
